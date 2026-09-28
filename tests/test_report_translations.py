"""
Calidad de traducción de TODOS los documentos generados (es / en / it).

Genera de extremo a extremo actas (Word y PDF), exportaciones CSV y etiquetas
QR en cada idioma y verifica que:
  · los rótulos clave del idioma solicitado están presentes;
  · no se filtran rótulos en español a los documentos en inglés/italiano;
  · los valores canónicos de catálogo (estados, tipos de movimiento, acciones
    de auditoría, estados de orden, sí/no) se traducen;
  · la fecha larga respeta el formato de cada idioma;
  · el español lleva tildes/ñ correctas.
"""
import csv
import io
import re
from datetime import date, datetime

import pytest
from docx import Document

from app.services.documents import DocumentService, _razon_social, _format_fecha_larga
from app.services.report_i18n import (
    format_amount,
    format_long_date,
    translate_catalog,
    yes_no,
)

LANGS = ("es", "en", "it")


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def _docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def _pdf_text(content: bytes) -> str:
    pypdf = pytest.importorskip("pypdf")
    reader = pypdf.PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() for page in reader.pages)


def _norm_ws(text: str) -> str:
    return re.sub(r"\s+", " ", text)


def _parse_csv(text: str):
    rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
    return rows[0], rows[1:]


def _has_word(text: str, word: str) -> bool:
    return re.search(rf"(?<!\w){re.escape(word)}(?!\w)", text) is not None


# Rótulos del acta que NO deben aparecer fuera del español.
# (En italiano "SERIE", "DISPOSITIVO", "MARCA" y "documento" son palabras propias.)
_SPANISH_ACTA_LEAKS_COMMON = [
    "Descripción", "Colaborador", "Departamento de TI", "Recibe", "Entrega",
    "Mensajero externo", "No pertenece", "Asignación", "Devolución", "septiembre",
    "MODELO", "equipo", "de la empresa", "respaldo", "administración",
]
_SPANISH_ACTA_LEAKS = {
    "en": _SPANISH_ACTA_LEAKS_COMMON + ["SERIE", "DISPOSITIVO", "MARCA", "documento", "NOTA"],
    "it": _SPANISH_ACTA_LEAKS_COMMON + ["Guatemala, 24 de"],
}

_EXPECTED_ACTA = {
    "es": {
        "comun": ["DISPOSITIVO", "MARCA", "MODELO", "SERIE", "Descripción:", "Colaborador",
                  "uso exclusivo de Lombardi", "Lombardi S.A."],
        "entrega": ["NOTA: ENTREGA DE EQUIPO", "Recibe el equipo", "Entrega y autoriza", "Departamento de TI",
                    "Asignación de equipo."],
        "descargo": ["NOTA: DEVOLUCIÓN A LOMBARDI", "Entrega el equipo", "Recibe",
                     "Devolución de equipo.", "Mensajero externo", "No pertenece a Lombardi"],
    },
    "en": {
        "comun": ["DEVICE", "BRAND", "MODEL", "SERIAL NUMBER", "Description:", "Employee",
                  "exclusive use of Lombardi", "Lombardi S.A."],
        "entrega": ["NOTE: EQUIPMENT HANDOVER", "Received by", "Delivered and authorized by", "IT Department",
                    "Equipment assignment."],
        "descargo": ["NOTE: RETURN OF EQUIPMENT TO LOMBARDI", "Returned by", "Received by",
                     "Equipment return.", "External courier", "Not affiliated with Lombardi"],
    },
    "it": {
        "comun": ["DISPOSITIVO", "MARCA", "MODELLO", "NUMERO DI SERIE", "Descrizione:",
                  "Collaboratore", "uso esclusivo di Lombardi", "Lombardi S.A."],
        "entrega": ["NOTA: CONSEGNA DI APPARECCHIATURE", "Riceve l'apparecchiatura",
                    "Consegna e autorizza", "Reparto IT", "Assegnazione di apparecchiature."],
        "descargo": ["NOTA: RESTITUZIONE DI APPARECCHIATURE A LOMBARDI",
                     "Restituisce l'apparecchiatura", "Riceve", "Restituzione di apparecchiature.",
                     "Corriere esterno", "Esterno a Lombardi"],
    },
}

_EXPECTED_DATE_FRAGMENT = {
    "es": re.compile(r"Guatemala, \d{1,2} de (enero|febrero|marzo|abril|mayo|junio|julio|agosto|"
                     r"septiembre|octubre|noviembre|diciembre) de \d{4}"),
    "en": re.compile(r"Guatemala, (January|February|March|April|May|June|July|August|September|"
                     r"October|November|December) \d{1,2}, \d{4}"),
    "it": re.compile(r"Guatemala, (\d{1,2}|1º) (gennaio|febbraio|marzo|aprile|maggio|giugno|luglio|"
                     r"agosto|settembre|ottobre|novembre|dicembre) \d{4}"),
}


async def _setup_movement(client, auth_headers, domain_seed, session):
    """Crea una asignación y configura la empresa 'Lombardi' (sin logo)."""
    from app.repositories.governance import GovernanceRepository

    d = domain_seed
    config = await GovernanceRepository(session).get_config()
    config.SYS_Nombre_Empresa = "Lombardi"
    config.SYS_Logo_URL = None
    await session.commit()

    r = await client.post(
        "/api/v1/trazabilidad/movimientos",
        json={"ACT_Activo": d["act_1"], "PER_Persona": d["alice"],
              "ARE_Area": d["area"], "TMO_Tipo_Movimiento": d["tmo_asg"]},
        headers=auth_headers,
    )
    assert r.status_code in (200, 201), r.text
    return r.json()["MOV_Movimiento"]


# ---------------------------------------------------------------------------
# Actas Word / PDF
# ---------------------------------------------------------------------------
@pytest.mark.asyncio
@pytest.mark.parametrize("lang", LANGS)
@pytest.mark.parametrize("formato", ("docx", "pdf"))
@pytest.mark.parametrize("tipo", ("entrega", "descargo"))
async def test_acta_traducida_completa(client, auth_headers, domain_seed, session, lang, formato, tipo):
    mid = await _setup_movement(client, auth_headers, domain_seed, session)
    params = {"formato": formato, "tipo": tipo, "lang": lang}
    if tipo == "descargo":
        params["mensajero"] = "Pedro Pérez"
    r = await client.get(f"/api/v1/trazabilidad/acta/{mid}", params=params, headers=auth_headers)
    assert r.status_code == 200, r.text

    text = _docx_text(r.content) if formato == "docx" else _pdf_text(r.content)
    flat = _norm_ws(text)

    for label in _EXPECTED_ACTA[lang]["comun"] + _EXPECTED_ACTA[lang][tipo]:
        assert label in flat, f"[{lang}/{tipo}/{formato}] falta el rótulo {label!r}"

    assert _EXPECTED_DATE_FRAGMENT[lang].search(flat), f"[{lang}] fecha mal formateada: {flat[:200]}"
    assert "Pedro Pérez" in flat or tipo == "entrega"  # acentos del usuario se conservan
    assert "Alice Test" in flat

    if lang != "es":
        for leak in _SPANISH_ACTA_LEAKS[lang]:
            assert not _has_word(flat, leak), f"[{lang}/{tipo}/{formato}] texto en español: {leak!r}"

    cd = r.headers["content-disposition"]
    expected_name = {
        ("es", "entrega"): "Acta_Entrega_", ("es", "descargo"): "Acta_Descargo_",
        ("en", "entrega"): "Handover_Record_", ("en", "descargo"): "Return_Record_",
        ("it", "entrega"): "Verbale_Consegna_", ("it", "descargo"): "Verbale_Restituzione_",
    }[(lang, tipo)]
    assert expected_name in cd and cd.endswith(f'.{formato}"')


@pytest.mark.asyncio
async def test_acta_lote_nombre_de_archivo_localizado(client, auth_headers, domain_seed, session):
    mid = await _setup_movement(client, auth_headers, domain_seed, session)
    r = await client.post(
        "/api/v1/trazabilidad/acta/lote",
        params={"formato": "pdf", "tipo": "entrega", "lang": "en"},
        json={"movimientos_ids": [mid]},
        headers=auth_headers,
    )
    assert r.status_code == 200, r.text
    assert "Handover_Record_Batch_" in r.headers["content-disposition"]


def test_acta_espanol_con_tildes_y_redaccion_formal():
    from app.services.documents import _ACTA_TEXTS_BY_LANG, _PIE_LEGAL_BY_LANG
    es = _ACTA_TEXTS_BY_LANG["es"]
    joined = " ".join(v if isinstance(v, str) else " ".join(v)
                      for t in es.values() for v in t.values()) + _PIE_LEGAL_BY_LANG["es"]
    # Concordancia corregida ("se detalla las características" era incorrecto).
    assert "se detalla las" not in joined
    assert "se detallan las características" in joined
    for palabra in ("características", "administración", "devolución", "electrónico", "informáticos"):
        assert palabra in joined.lower()


def test_acta_pdf_escapa_datos_del_usuario():
    """Nombres con '&' o '<' no deben romper el PDF (reportlab usa mini-HTML)."""
    from app.services.documents import _ACTA_TEXTS_BY_LANG, _LABELS_BY_LANG, _PIE_LEGAL_BY_LANG
    data = {
        "tipo": "entrega", "lang": "en", "txt": _ACTA_TEXTS_BY_LANG["en"]["entrega"],
        "labels": _LABELS_BY_LANG["en"], "pie_legal": _PIE_LEGAL_BY_LANG["en"],
        "empresa": "Smith & Sons", "logo": None, "primary": (31, 58, 95),
        "colaborador": "Ana <Admin> López", "mensajero": None, "ciudad": "Guatemala",
        "codigo_form": "F.IT.GUA.04.01", "fecha_larga": "September 24, 2026",
        "acta_ref": "AB12CD34", "motivo": "Replacement: <old unit> & charger",
        "items": [{"dispositivo": "Laptop", "marca": "DELL", "modelo": "5440", "serie": "X&Y"}],
    }
    svc = DocumentService.__new__(DocumentService)
    text = _norm_ws(_pdf_text(svc._render_pdf(data).getvalue()))
    assert "Smith & Sons" in text
    assert "Ana <Admin> López" in text
    assert "<old unit> & charger" in text
    assert _svc_docx_ok(svc, data)


def _svc_docx_ok(svc, data) -> bool:
    return svc._render_docx(data).getvalue()[:2] == b"PK"


def test_razon_social_no_duplica_forma_societaria():
    assert _razon_social("Lombardi") == "Lombardi S.A."
    assert _razon_social("Lombardi S.A.") == "Lombardi S.A."
    assert _razon_social("Lombardi, S.A.") == "Lombardi, S.A."
    assert _razon_social("Acme Inc.") == "Acme Inc."


@pytest.mark.parametrize("lang,expected", [
    ("es", "24 de septiembre de 2026"),
    ("en", "September 24, 2026"),
    ("it", "24 settembre 2026"),
])
def test_fecha_larga_por_idioma(lang, expected):
    assert format_long_date(date(2026, 9, 24), lang) == expected
    assert _format_fecha_larga(datetime(2026, 9, 24, 10, 0), lang) == expected


def test_fecha_larga_italiano_primer_dia_ordinal():
    assert format_long_date(date(2026, 3, 1), "it") == "1º marzo 2026"
    assert format_long_date(date(2026, 3, 1), "es") == "1 de marzo de 2026"


# ---------------------------------------------------------------------------
# Catálogos canónicos
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("catalog,value,expected", [
    ("estado_operativo", "Asignado", {"es": "Asignado", "en": "Assigned", "it": "Assegnato"}),
    ("estado_operativo", "Disponible", {"es": "Disponible", "en": "Available", "it": "Disponibile"}),
    ("estado_operativo", "En Bodega", {"es": "En Bodega", "en": "In Storage", "it": "In magazzino"}),
    ("estado_operativo", "En Reparación", {"es": "En Reparación", "en": "Under Repair", "it": "In riparazione"}),
    ("estado_operativo", "Baja", {"es": "Baja", "en": "Decommissioned", "it": "Dismesso"}),
    ("estado_operativo", "STATUS.ASSIGNED", {"es": "Asignado", "en": "Assigned", "it": "Assegnato"}),
    ("tipo_movimiento", "Asignación", {"es": "Asignación", "en": "Assignment", "it": "Assegnazione"}),
    ("tipo_movimiento", "Devolucion", {"es": "Devolución", "en": "Return", "it": "Restituzione"}),
    ("tipo_movimiento", "Transferencia", {"es": "Transferencia", "en": "Transfer", "it": "Trasferimento"}),
    ("tipo_movimiento", "Préstamo", {"es": "Préstamo", "en": "Loan", "it": "Prestito"}),
    ("tipo_movimiento", "Ingreso", {"es": "Ingreso", "en": "Intake", "it": "Presa in carico"}),
    ("estado_orden", "BORRADOR", {"es": "Borrador", "en": "Draft", "it": "Bozza"}),
    ("estado_orden", "RECIBIDA", {"es": "Recibida", "en": "Received", "it": "Ricevuto"}),
    ("accion_auditoria", "LOGIN_SUCCESS",
     {"es": "Inicio de sesión exitoso", "en": "Successful sign-in", "it": "Accesso riuscito"}),
    ("tipo_activo", "Impresora", {"es": "Impresora", "en": "Printer", "it": "Stampante"}),
    ("unidad", "unidad", {"es": "unidad", "en": "unit", "it": "unità"}),
])
def test_traduccion_catalogos_canonicos(catalog, value, expected):
    for lang in LANGS:
        assert translate_catalog(catalog, value, lang) == expected[lang]


def test_valores_libres_no_se_alteran():
    assert translate_catalog("estado_operativo", "Estado personalizado XYZ", "en") == "Estado personalizado XYZ"
    assert translate_catalog("tipo_activo", "", "it") == ""
    assert translate_catalog("tipo_activo", None, "it") == ""


def test_si_no_y_montos():
    assert [yes_no(True, lang) for lang in LANGS] == ["Sí", "Yes", "Sì"]
    assert [yes_no(False, lang) for lang in LANGS] == ["No", "No", "No"]
    assert format_amount(1500) == "1500.00"
    assert format_amount("12.5") == "12.50"
    assert format_amount(None) == ""


# ---------------------------------------------------------------------------
# Exportaciones CSV
# ---------------------------------------------------------------------------
_SPANISH_CSV_LEAKS = [
    "Código", "Codigo", "Serie", "Fecha", "Estado", "Tipo_", "Tipo de", "Asignado", "Asignación",
    "Disponible", "En Bodega", "Devolución", "Correo", "Proveedor", "Borrador", "Existencia",
    "Teléfono", "Dirección", "Acción", "Accion", "Sí", "SI", "unidad", "Observac", "Número",
]


async def _seed_reportes(session):
    from app.models.consumable import Consumible
    from app.models.procurement import OrdenCompra, Proveedor

    prv = Proveedor(PRV_Nombre="Proveedor Uno", PRV_Activo=True)
    session.add(prv)
    await session.flush()
    session.add_all([
        Consumible(CON_Nombre="Toner HP 26A", CON_Categoria="Toner", CON_Unidad="unidad",
                   CON_Stock_Actual=1, CON_Stock_Minimo=3, CON_Activo=True),
        OrdenCompra(OCO_Numero="OC-0001", OCO_Fecha=date(2026, 9, 1), OCO_Estado="BORRADOR",
                    OCO_Moneda="USD", OCO_Total=1250, PRV_Proveedor=prv.PRV_Proveedor),
    ])
    await session.commit()


_CSV_EXPECTED = {
    "activos": {
        "es": ["Código interno", "Número de serie", "Estado operativo", "Vencimiento de garantía"],
        "en": ["Internal code", "Serial number", "Operational status", "Warranty expiration"],
        "it": ["Codice interno", "Numero di serie", "Stato operativo", "Scadenza garanzia"],
    },
    "movimientos": {
        "es": ["Fecha de asignación", "Fecha de devolución", "Área", "Tipo de movimiento", "Observaciones"],
        "en": ["Assignment date", "Return date", "Area", "Movement type", "Notes"],
        "it": ["Data di assegnazione", "Data di restituzione", "Area", "Tipo di movimento", "Note"],
    },
    "consumibles": {
        "es": ["Categoría", "Existencia actual", "Existencia mínima", "Bajo el mínimo"],
        "en": ["Category", "Current stock", "Minimum stock", "Below minimum"],
        "it": ["Categoria", "Giacenza attuale", "Giacenza minima", "Sotto la scorta minima"],
    },
    "proveedores": {
        "es": ["Identificación fiscal", "Correo electrónico", "Teléfono", "Dirección"],
        "en": ["Tax ID", "Email", "Phone", "Address"],
        "it": ["Identificativo fiscale", "Email", "Telefono", "Indirizzo"],
    },
    "ordenes": {
        "es": ["Número de orden", "Proveedor", "Estado", "Moneda"],
        "en": ["Order number", "Supplier", "Status", "Currency"],
        "it": ["Numero ordine", "Fornitore", "Stato", "Valuta"],
    },
    "auditoria": {
        "es": ["Fecha y hora", "Acción", "Entidad afectada", "IP de origen", "Detalle (JSON)"],
        "en": ["Date and time", "Action", "Affected entity", "Source IP", "Details (JSON)"],
        "it": ["Data e ora", "Azione", "Entità interessata", "IP di origine", "Dettagli (JSON)"],
    },
}

_CSV_FILENAME_PREFIX = {
    "activos": {"es": "activos_", "en": "assets_", "it": "asset_"},
    "movimientos": {"es": "movimientos_", "en": "movements_", "it": "movimenti_"},
    "consumibles": {"es": "consumibles_", "en": "consumables_", "it": "materiali_di_consumo_"},
    "proveedores": {"es": "proveedores_", "en": "suppliers_", "it": "fornitori_"},
    "ordenes": {"es": "ordenes_de_compra_", "en": "purchase_orders_", "it": "ordini_di_acquisto_"},
    "auditoria": {"es": "auditoria_", "en": "audit_log_", "it": "registro_audit_"},
}


@pytest.mark.asyncio
@pytest.mark.parametrize("lang", LANGS)
async def test_exports_csv_traducidos(client, auth_headers, domain_seed, session, lang):
    await _setup_movement(client, auth_headers, domain_seed, session)
    await _seed_reportes(session)

    for name, expected in _CSV_EXPECTED.items():
        r = await client.get(f"/api/v1/export/{name}.csv", params={"lang": lang}, headers=auth_headers)
        assert r.status_code == 200, (name, r.text)
        headers, rows = _parse_csv(r.text)
        for label in expected[lang]:
            assert label in headers, f"[{name}/{lang}] falta el encabezado {label!r}: {headers}"
        assert not any("_" in h for h in headers), f"[{name}/{lang}] encabezado técnico: {headers}"
        assert re.search(rf'filename="{_CSV_FILENAME_PREFIX[name][lang]}\d{{8}}_\d{{6}}\.csv"',
                         r.headers["content-disposition"])

        if lang != "es" and name != "auditoria":
            # (auditoría: el snapshot JSON es dato forense y se conserva tal cual)
            body = "\n".join(",".join(row) for row in [headers, *rows])
            body = body.replace("Proveedor Uno", "")  # dato libre del usuario
            for leak in _SPANISH_CSV_LEAKS:
                assert not _has_word(body, leak), f"[{name}/{lang}] texto en español: {leak!r}"
        if lang != "es" and name == "auditoria":
            for leak in ("Código", "Acción", "Accion", "Fecha"):
                assert not any(_has_word(h, leak) for h in headers)


@pytest.mark.asyncio
@pytest.mark.parametrize("lang,estado,tipo_mov,si,orden,unidad", [
    ("es", "Asignado", "Asignación", "Sí", "Borrador", "unidad"),
    ("en", "Assigned", "Assignment", "Yes", "Draft", "unit"),
    ("it", "Assegnato", "Assegnazione", "Sì", "Bozza", "unità"),
])
async def test_exports_csv_valores_de_catalogo(client, auth_headers, domain_seed, session,
                                              lang, estado, tipo_mov, si, orden, unidad):
    await _setup_movement(client, auth_headers, domain_seed, session)
    await _seed_reportes(session)

    _, rows = _parse_csv((await client.get("/api/v1/export/activos.csv", params={"lang": lang},
                                           headers=auth_headers)).text)
    lap1 = next(r for r in rows if r[0] == "LAP-001")
    assert lap1[6] == estado
    assert lap1[7] == "2024-01-01"  # ISO 8601: inequívoco en cualquier Excel

    _, rows = _parse_csv((await client.get("/api/v1/export/movimientos.csv", params={"lang": lang},
                                           headers=auth_headers)).text)
    assert rows[0][7] == tipo_mov
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", rows[0][0])

    _, rows = _parse_csv((await client.get("/api/v1/export/consumibles.csv", params={"lang": lang},
                                           headers=auth_headers)).text)
    assert rows[0][2] == unidad
    assert rows[0][5] == si and rows[0][6] == si

    _, rows = _parse_csv((await client.get("/api/v1/export/ordenes.csv", params={"lang": lang},
                                           headers=auth_headers)).text)
    assert rows[0][3] == orden
    assert rows[0][5] == "1250.00"
    assert rows[0][2] == "2026-09-01"


# ---------------------------------------------------------------------------
# Etiquetas QR
# ---------------------------------------------------------------------------
@pytest.mark.asyncio
@pytest.mark.parametrize("lang,prefix,filename", [
    ("es", "N.º de serie: SER-001", "etiquetas-activos.pdf"),
    ("en", "Serial no.: SER-001", "asset-labels.pdf"),
    ("it", "N. di serie: SER-001", "etichette-asset.pdf"),
])
async def test_etiquetas_traducidas(client, auth_headers, domain_seed, lang, prefix, filename):
    d = domain_seed
    r = await client.post(
        "/api/v1/core/activos/etiquetas.pdf",
        params={"lang": lang},
        json={"activos_ids": [d["act_1"], d["act_2"]]},
        headers=auth_headers,
    )
    assert r.status_code == 200, r.text
    assert filename in r.headers["content-disposition"]
    text = _pdf_text(r.content)
    assert prefix in text
    if lang != "es":
        assert "Serie:" not in text


@pytest.mark.asyncio
async def test_etiquetas_sin_lang_mantienen_espanol(client, auth_headers, domain_seed):
    """Compatibilidad: sin ?lang el endpoint sigue respondiendo en español."""
    d = domain_seed
    r = await client.post("/api/v1/core/activos/etiquetas.pdf",
                          json={"activos_ids": [d["act_1"]]}, headers=auth_headers)
    assert r.status_code == 200, r.text
    assert "N.º de serie: SER-001" in _pdf_text(r.content)
